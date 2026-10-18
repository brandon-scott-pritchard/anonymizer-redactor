"""The final check's flagged pages, fixed on the review screen."""

import zipfile
from types import SimpleNamespace

import docx
import pytest

from redactor import final_review, grouping, mapping, pipeline, review
from redactor.engine import Settings

PASSWORD = "test-only-final-review"         # a throwaway key password for these tests


def _delivered(archive, member):
    import io
    with zipfile.ZipFile(archive) as zf:
        document = docx.Document(io.BytesIO(zf.read(member)))
    return "\n".join(p.text for p in document.paragraphs)


@pytest.fixture
def ran(tmp_path):
    """A two-file set run once, with a name and a number the run missed."""
    first = tmp_path / "letter.docx"
    d = docx.Document()
    d.add_paragraph("Dear Abigail N Okafor,")
    d.add_paragraph("Your neighbour Bgail Okfr Smyth signed for it. Locker 88213377.")
    d.save(first)
    second = tmp_path / "note.docx"
    d = docx.Document()
    d.add_paragraph("Locker 88213377 is due. Bgail Okfr Smyth again.")
    d.save(second)
    store, _ = review.build_store([("Abigail N Okafor", "person")])
    settings = Settings(docx_mode="anonymize")
    pipeline.prescan([first, second], store, settings)
    out = tmp_path / "out"
    out.mkdir()
    result = pipeline.run_job([first, second], store, settings, out, PASSWORD)
    members = sorted(o.delivered_name for o in result.outcomes)
    doc_set = SimpleNamespace(id="set", run_result=result, store=store, run_settings=settings,
                              run_password=PASSWORD, audit={},
                              final_flags=[{"file": m, "page": 1, "kev": 0.9, "matched": []}
                                           for m in members])
    return doc_set, members


def test_the_run_missed_what_the_fix_will_take(ran):
    doc_set, members = ran
    text = _delivered(doc_set.run_result.archive, members[0])
    assert "Abigail" not in text                       # the ticked name went
    assert "Bgail Okfr Smyth" in text and "88213377" in text


def test_a_chosen_name_changes_this_file_and_lists_the_others_holding_it(ran, monkeypatch):
    """One file at a time: the page's file is made again, and the other file
    that still shows the same text comes onto the list - unchanged."""
    doc_set, members = ran
    monkeypatch.setenv("REDACTOR_KEV", "0")            # the matcher reads it back alone
    doc_set.final_flags = [doc_set.final_flags[0]]     # only the first file was flagged
    done = final_review.fix(doc_set, members[0], 1, "Bgail Okfr Smyth", "person", "")
    assert done.get("error") is None and done["page_clean"]
    assert "Bgail Okfr Smyth" not in _delivered(doc_set.run_result.archive, members[0])
    assert "Bgail Okfr Smyth" in _delivered(doc_set.run_result.archive, members[1])
    assert done["added"] == 1
    assert [(f["file"], f["matched"]) for f in doc_set.final_flags] == [
        (members[1], ["Bgail Okfr Smyth"])]


def test_a_chosen_number_takes_a_placeholder_and_the_key_is_written_again(ran, monkeypatch):
    doc_set, members = ran
    monkeypatch.setenv("REDACTOR_KEV", "0")
    final_review.fix(doc_set, members[0], 1, "88213377", "bank_account", "")
    assert "88213377" not in _delivered(doc_set.run_result.archive, members[0])
    key = mapping.read_encrypted_key(doc_set.run_result.key_path, PASSWORD)
    assert "88213377" in str(key)


def test_the_preview_says_what_a_choice_will_read_as_and_changes_nothing(ran):
    doc_set, members = ran
    before = len(doc_set.store.entities)
    seen = final_review.preview(doc_set, members[0], "88213377", "bank_account", "")
    assert seen["reads_as"].startswith("[") and not seen["short"]
    assert len(doc_set.store.entities) == before


def test_part_of_a_ticked_name_given_no_row_is_warned_about(ran):
    doc_set, members = ran
    seen = final_review.preview(doc_set, members[0], "Abigail", "person", "")
    assert any("part of Abigail N Okafor" in w for w in seen.get("warnings", []))


def test_an_initial_is_never_registered_across_the_set(ran):
    """"C" as a form read as a surname, took a whole invented name, matched no
    lone C, and stands in every document: on its own it is replaced where it
    stands, on a PDF page, as the person's invented initial."""
    doc_set, members = ran
    head = next(e for e in doc_set.store.entities.values() if e.canonical == "Abigail N Okafor")
    docx_page = final_review.preview(doc_set, members[0], "A", "person", head.key)
    assert docx_page.get("error") and "PDF" in docx_page["error"]
    assert final_review.preview(doc_set, "page.pdf", "A.", "person", head.key)["blacked_out"]
    doc_set.run_settings.pdf_mode = "anonymize"
    seen = final_review.preview(doc_set, "page.pdf", "A.", "person", head.key)
    assert seen["short"] and seen["reads_as"] == head.surrogate.first[:1] + "."
    nobody = final_review.preview(doc_set, "page.pdf", "A", "person", "")
    assert nobody.get("error") and "Whose" in nobody["error"]


def test_an_initial_on_a_pdf_page_is_covered_where_it_stands(tmp_path, monkeypatch):
    import pymupdf

    from redactor import final_review as F

    monkeypatch.setenv("REDACTOR_KEV", "0")
    pdf = tmp_path / "statement.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((60, 80), "Account holder: Abigail N Okafor", fontsize=11)
    page.insert_text((60, 110), "Joint owner: C", fontsize=11)
    page.insert_text((60, 140), "Schedule C attached", fontsize=11)
    doc.save(pdf)
    store, _ = review.build_store([("Abigail N Okafor", "person"), ("Cody J Okafor", "person")])
    settings = Settings(pdf_mode="anonymize")
    pipeline.prescan([pdf], store, settings)
    out = tmp_path / "out"
    out.mkdir()
    result = pipeline.run_job([pdf], store, settings, out, PASSWORD)
    member = result.outcomes[0].delivered_name
    doc_set = SimpleNamespace(id="s", run_result=result, store=store, run_settings=settings,
                              run_password=PASSWORD, audit={},
                              final_flags=[{"file": member, "page": 1, "kev": 0.7, "matched": []}])
    words = F.page_words(result.archive, member, 1)["words"]
    lone = [w for w in words if w["text"] == "C"]
    assert len(lone) == 2                               # the initial and "Schedule C"
    initial = min(lone, key=lambda w: w["y0"])
    cody = next(e for e in store.entities.values() if e.canonical == "Cody J Okafor")
    done = F.fix(doc_set, member, 1, "C", "person", cody.key, [initial["i"]], ["C"])
    assert done.get("error") is None, done
    with zipfile.ZipFile(result.archive) as zf:
        with pymupdf.open(stream=zf.read(member), filetype="pdf") as d:
            text = d[0].get_text()
            after = [w for w in d[0].get_text("words")]
    assert "Joint owner: C" not in text
    assert "Schedule C attached" in text                # nothing else was touched
    # where the initial stood, the invented one
    x0, y0, x1, y1 = (initial[k] / F.ZOOM for k in ("x0", "y0", "x1", "y1"))
    there = [w[4] for w in after if w[0] < x1 + 2 and w[2] > x0 - 2 and w[1] < y1 and w[3] > y0]
    assert there == [cody.surrogate.first[:1]]


def test_a_form_of_a_ticked_person_takes_their_invented_name(ran, monkeypatch):
    doc_set, members = ran
    monkeypatch.setenv("REDACTOR_KEV", "0")
    head = next(e for e in doc_set.store.entities.values() if e.canonical == "Abigail N Okafor")
    final_review.fix(doc_set, members[0], 1, "Bgail Okfr", "person", head.key)
    form = doc_set.store.get("person", "Bgail Okfr") or next(
        e for e in doc_set.store.entities.values() if e.canonical == "Bgail Okfr")
    assert grouping.head_of(doc_set.store, form) is head
    assert form.replacement.split()[-1] == head.surrogate.last


def test_a_page_looked_at_with_nothing_on_it_comes_off_the_list(ran):
    doc_set, members = ran
    assert final_review.dismiss(doc_set, members[0], 1) == 1
    assert [f["file"] for f in doc_set.final_flags] == [members[1]]


def test_suspects_mark_a_word_a_letter_or_two_off_a_ticked_name():
    store, _ = review.build_store([("Abigail N Okafor", "person")])
    words = [{"i": i, "text": t} for i, t in enumerate(
        ["Paid", "to", "ABIGAIL", "OKAF0R", "for", "toABIGAILOKAFOR", "groceries"])]
    marked = final_review.suspects(words, store, [])["words"]
    assert 3 in marked and 5 in marked and 6 not in marked and 0 not in marked


def test_suspects_mark_a_ticked_first_names_nickname():
    store, _ = review.build_store([("William J Harrington", "person")])
    words = [{"i": i, "text": t} for i, t in enumerate(["Transfer", "to", "Billy", "Harringtn"])]
    marked = final_review.suspects(words, store, [])["words"]
    assert marked == [2, 3]


def test_a_kind_left_as_written_is_refused_not_reported_clean(ran, monkeypatch):
    """A loan number filed as a loan number is never replaced; saying the page
    reads clean would be false."""
    doc_set, members = ran
    monkeypatch.setenv("REDACTOR_KEV", "0")
    done = final_review.fix(doc_set, members[0], 1, "88213377", "loan_number", "")
    assert done.get("error") and "left as written" in done["error"]
    assert "88213377" in _delivered(doc_set.run_result.archive, members[0])
    assert "loan_number" not in {k["key"] for k in final_review.kinds(doc_set.run_settings)}


def test_a_file_that_cannot_be_made_again_is_never_reported_clean(ran, monkeypatch):
    doc_set, members = ran
    monkeypatch.setenv("REDACTOR_KEV", "0")
    for outcome in doc_set.run_result.outcomes:
        outcome.source.unlink()                      # the original is gone
    done = final_review.fix(doc_set, members[0], 1, "88213377", "bank_account", "")
    assert not done["page_clean"] and "could not be made again" in done["how"]
    assert any(f["file"] == members[0] for f in doc_set.final_flags)


def test_text_still_on_the_page_keeps_the_page_on_the_list(ran, monkeypatch):
    """A file made again that still holds what was chosen - a value the
    literal pass will not take, here stood in for by a remake that changes
    nothing - is not clean, and the page stays listed."""
    from redactor import audit

    doc_set, members = ran
    monkeypatch.setenv("REDACTOR_KEV", "0")
    monkeypatch.setattr(audit, "remake", lambda *a, **k: {})
    done = final_review.fix(doc_set, members[0], 1, "Locker", "bank_account", "")
    assert not done["page_clean"] and "still on the page" in done["how"]
    assert any(f["file"] == members[0] and f["page"] == 1 for f in doc_set.final_flags)


def test_kev_unsure_pages_are_listed_after_the_flagged_ones(tmp_path):
    from redactor import audit

    check = audit.Check(kev_note="Kev", pages=3)
    check.flags = [audit.Flag("a.pdf", 1, 0.8)]
    check.doubtful = [audit.Flag("b.pdf", 2, 0.35, doubtful=True)]
    summary = check.summary()
    assert [f["unsure"] for f in summary["flags"]] == [False, True]
    check.path = tmp_path / "final-check.txt"
    text = audit.render(check)
    assert "Kev unsure - 1 page(s)" in text and "b.pdf, page 2" in text
